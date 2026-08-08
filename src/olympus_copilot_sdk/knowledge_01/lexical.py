from __future__ import annotations

import csv
import json
import math
import re
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Final

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from olympus_copilot_sdk.knowledge_01.retrieval import IndexSummary, SearchResult

TEXT_SUFFIXES: Final = {
    ".css",
    ".js",
    ".json",
    ".md",
    ".py",
    ".rst",
    ".toml",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
MAX_FILE_CHARACTERS: Final = 500_000
CHUNK_CHARACTERS: Final = 3_000
CHUNK_OVERLAP: Final = 300


@dataclass(frozen=True)
class Chunk:
    source: str
    text: str
    terms: frozenset[str]


class KnowledgeBase:
    def __init__(self, data_directory: Path) -> None:
        self.data_directory = data_directory.resolve()
        self._chunks: list[Chunk] = []
        indexed: list[str] = []
        skipped: list[str] = []
        for source, text in iter_documents(self.data_directory, skipped):
            indexed.append(source)
            self._chunks.extend(_make_chunks(source, text[:MAX_FILE_CHARACTERS]))
        self.summary = IndexSummary(tuple(indexed), tuple(skipped), len(self._chunks))
        self._document_frequency = Counter(term for chunk in self._chunks for term in chunk.terms)

    def search(self, query: str, limit: int = 6) -> list[SearchResult]:
        query_terms = _terms(query)
        if not query_terms:
            return []
        ranked: list[SearchResult] = []
        query_identifiers = _identifiers(query)
        query_weight = sum(self._idf(term) for term in query_terms)
        for chunk in self._chunks:
            overlap = query_terms & chunk.terms
            if not overlap:
                continue
            lexical_score = (
                sum(
                    self._idf(term) * (2.0 if term in chunk.source.lower() else 1.0)
                    for term in overlap
                )
                / query_weight
            )
            query_norm = math.sqrt(sum(self._idf(term) ** 2 for term in query_terms))
            chunk_norm = math.sqrt(sum(self._idf(term) ** 2 for term in chunk.terms))
            vector_score = (
                sum(self._idf(term) ** 2 for term in overlap) / (query_norm * chunk_norm)
                if query_norm and chunk_norm
                else 0.0
            )
            score = (
                lexical_score
                + vector_score
                + 4.0 * len(query_identifiers & _identifiers(chunk.text))
            )
            if Path(chunk.source).suffix.lower() in {".css", ".js"}:
                score *= 0.1
            ranked.append(SearchResult(chunk.source, chunk.text, score))
        ranked.sort(key=lambda result: result.score, reverse=True)
        return ranked[:limit]

    def _idf(self, term: str) -> float:
        return math.log((len(self._chunks) + 1) / (self._document_frequency.get(term, 0) + 1)) + 1.0


def iter_documents(
    data_directory: Path,
    skipped: list[str] | None = None,
) -> Iterator[tuple[str, str]]:
    if not data_directory.exists():
        return
    for path in sorted(data_directory.rglob("*")):
        if not path.is_file() or path.name.startswith("."):
            continue
        source = str(path.relative_to(data_directory))
        try:
            text = _extract_text(path)
        except (OSError, ValueError, TypeError, KeyError):
            if skipped is not None:
                skipped.append(source)
            continue
        if not text.strip():
            if skipped is not None:
                skipped.append(source)
            continue
        yield source, text


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
        for element in soup(["script", "style"]):
            element.decompose()
        return soup.get_text(" ", strip=True)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join([*paragraphs, *table_rows])
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                rows.append(f"Sheet: {worksheet.title}")
                rows.extend(
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in worksheet.iter_rows(values_only=True)
                )
        finally:
            workbook.close()
        return "\n".join(rows)
    if suffix == ".csv":
        with path.open(encoding="utf-8", errors="replace", newline="") as stream:
            return "\n".join(" | ".join(row) for row in csv.reader(stream))
    if suffix in TEXT_SUFFIXES or path.name.lower() in {"makefile", "dockerfile"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".json":
            return json.dumps(json.loads(text), indent=2, ensure_ascii=True)
        return text
    return ""


def _make_chunks(source: str, text: str) -> list[Chunk]:
    normalized = re.sub(r"\s+", " ", text).strip()
    chunks: list[Chunk] = []
    start = 0
    while start < len(normalized):
        end = min(start + CHUNK_CHARACTERS, len(normalized))
        chunk_text = normalized[start:end]
        chunks.append(Chunk(source, chunk_text, frozenset(_terms(chunk_text))))
        if end == len(normalized):
            break
        start = end - CHUNK_OVERLAP
    return chunks


def _terms(text: str) -> set[str]:
    return {
        _normalize_term(term)
        for term in re.findall(r"[a-z0-9]{2,}", text.lower())
        if term not in _STOPWORDS
    }


def _normalize_term(term: str) -> str:
    if len(term) > 4 and term.endswith("ies"):
        return f"{term[:-3]}y"
    if len(term) > 3 and term.endswith("s") and not term.endswith(("is", "ss", "us")):
        return term[:-1]
    return term


def _identifiers(text: str) -> set[str]:
    return {
        re.sub(r"\s+", " ", identifier.lower()).strip()
        for identifier in re.findall(
            r"\b(?:epic|feature|issue|story|task|ticket)\s*[-#:]?\s*\d+\b",
            text,
            flags=re.IGNORECASE,
        )
    }


_STOPWORDS: Final = {
    "a",
    "all",
    "an",
    "and",
    "about",
    "are",
    "as",
    "at",
    "be",
    "by",
    "can",
    "could",
    "describe",
    "detail",
    "details",
    "explain",
    "for",
    "from",
    "give",
    "how",
    "in",
    "is",
    "it",
    "know",
    "me",
    "of",
    "on",
    "or",
    "please",
    "list",
    "show",
    "specifically",
    "that",
    "the",
    "this",
    "tell",
    "to",
    "was",
    "what",
    "when",
    "where",
    "which",
    "who",
    "with",
    "you",
}
