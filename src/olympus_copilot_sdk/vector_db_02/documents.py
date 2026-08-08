from __future__ import annotations

import csv
import json
from collections.abc import Iterator
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

TEXT_SUFFIXES = {
    ".css",
    ".js",
    ".json",
    ".md",
    ".py",
    ".rst",
    ".toml",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}


def iter_documents(
    data_directory: Path,
    skipped: list[str] | None = None,
) -> Iterator[tuple[str, str]]:
    if not data_directory.exists():
        return
    for path in sorted(data_directory.rglob("*")):
        if not path.is_file() or path.name.startswith("."):
            continue
        source = str(path.relative_to(data_directory))
        try:
            text = extract_text(path)
        except (OSError, ValueError, TypeError, KeyError):
            if skipped is not None:
                skipped.append(source)
            continue
        if text.strip():
            yield source, text
        elif skipped is not None:
            skipped.append(source)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
        for element in soup(["script", "style"]):
            element.decompose()
        return soup.get_text(" ", strip=True)
    if suffix == ".pdf":
        return "\n\n".join(
            f"Page {page_number}\n{page.extract_text() or ''}"
            for page_number, page in enumerate(PdfReader(path).pages, start=1)
        )
    if suffix == ".docx":
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        rows = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join([*paragraphs, *rows])
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                rows.append(f"Sheet: {worksheet.title}")
                rows.extend(
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in worksheet.iter_rows(values_only=True)
                )
        finally:
            workbook.close()
        return "\n".join(rows)
    if suffix == ".csv":
        with path.open(encoding="utf-8", errors="replace", newline="") as stream:
            return "\n".join(" | ".join(row) for row in csv.reader(stream))
    if suffix in TEXT_SUFFIXES or path.name.lower() in {"makefile", "dockerfile"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return (
            json.dumps(json.loads(text), indent=2, ensure_ascii=True) if suffix == ".json" else text
        )
    return ""
