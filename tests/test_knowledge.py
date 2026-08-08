from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from olympus_copilot_sdk.knowledge_01.lexical import KnowledgeBase


def test_indexes_html_and_plain_text(tmp_path: Path) -> None:
    (tmp_path / "story.html").write_text(
        "<html><style>hidden</style><body>Toad visited the river bank.</body></html>"
    )
    (tmp_path / "notes.md").write_text("Mole packed a picnic basket.")

    knowledge = KnowledgeBase(tmp_path)

    assert knowledge.summary.indexed_files == ("notes.md", "story.html")
    results = knowledge.search("Who visited the river bank?")
    assert results[0].source == "story.html"
    assert "hidden" not in results[0].text


def test_skips_unsupported_files(tmp_path: Path) -> None:
    (tmp_path / "cover.jpg").write_bytes(b"not an image")

    knowledge = KnowledgeBase(tmp_path)

    assert knowledge.summary.skipped_files == ("cover.jpg",)
    assert knowledge.search("anything") == []


def test_skips_pdf_without_extractable_text(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with (tmp_path / "scan.pdf").open("wb") as stream:
        writer.write(stream)

    knowledge = KnowledgeBase(tmp_path)

    assert knowledge.summary.indexed_files == ()
    assert knowledge.summary.skipped_files == ("scan.pdf",)


def test_matches_plural_query_to_singular_document_term(tmp_path: Path) -> None:
    (tmp_path / "backlog.md").write_text("Epic 1: Platform Foundation")
    (tmp_path / "story.md").write_text("Show all the creatures in the literary story.")

    results = KnowledgeBase(tmp_path).search("Show all epics")

    assert results[0].source == "backlog.md"


def test_prioritizes_exact_epic_identifier_over_generic_language(tmp_path: Path) -> None:
    (tmp_path / "backlog.md").write_text(
        "Epic 3: Shared Framework. Epic 4: Specialist Agents and business workflows."
    )
    (tmp_path / "story.md").write_text(
        "Can you tell me about this story specifically? It has many details to explain."
    )

    results = KnowledgeBase(tmp_path).search("Can you tell me about Epic 4 specifically?")

    assert results[0].source == "backlog.md"
    assert "Epic 4" in results[0].text


def test_extracts_docx_tables(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Delivery backlog")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Epic 12"
    table.rows[0].cells[1].text = "Target release Q4"
    document.save(str(tmp_path / "backlog.docx"))

    results = KnowledgeBase(tmp_path).search("target release for epic 12")

    assert results[0].source == "backlog.docx"
    assert "Epic 12 | Target release Q4" in results[0].text
